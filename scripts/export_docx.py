#!/usr/bin/env python3
"""Export this repository's Markdown CV to an editable Word document.

Requires Python 3.10+ and requirements-docx.txt. No Word, Jekyll, or network
connection is needed to generate the file. Content always comes from Markdown.
"""

from __future__ import annotations

import argparse
from html.parser import HTMLParser
from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.shared import Inches, Pt, RGBColor
from markdown_it import MarkdownIt


ROOT = Path(__file__).resolve().parents[1]
MARKDOWN = MarkdownIt("commonmark", {"html": True})


class ContactParser(HTMLParser):
    """Keep visible contact text and links, leaving icon-only HTML behind."""

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.href = None

    def handle_starttag(self, tag, attrs):
        if tag not in {"div", "i", "a", "br"}:
            raise ValueError(f"Unsupported contact HTML: <{tag}>")
        if tag == "a":
            self.href = dict(attrs).get("href")
        elif tag == "br":
            self.parts.append((" ", None))

    def handle_endtag(self, tag):
        if tag == "a":
            self.href = None

    def handle_data(self, data):
        self.parts.append((re.sub(r"\s+", " ", data), self.href))


def add_text(paragraph, text, *, bold=False, italic=False, href=None):
    run = paragraph.add_run(text)
    run.bold = True if bold else None
    run.italic = True if italic else None
    if href:
        if not re.match(r"^(https?://|mailto:|tel:)", href):
            raise ValueError(f"Unsupported link target: {href}")
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), paragraph.part.relate_to(href, RT.HYPERLINK, is_external=True))
        link.append(run._r)
        paragraph._p.append(link)
    return run


def add_inline(paragraph, children):
    bold = italic = False
    href = None
    for token in children or []:
        kind = token.type
        if kind == "strong_open":
            bold = True
        elif kind == "strong_close":
            bold = False
        elif kind == "em_open":
            italic = True
        elif kind == "em_close":
            italic = False
        elif kind == "link_open":
            href = token.attrGet("href")
        elif kind == "link_close":
            href = None
        elif kind in {"text", "code_inline"}:
            run = add_text(paragraph, token.content, bold=bold,
                           italic=italic or kind == "code_inline", href=href)
            if kind == "code_inline":
                run.font.size = Pt(10)
        elif kind == "softbreak":
            add_text(paragraph, " ", bold=bold, italic=italic, href=href)
        elif kind == "hardbreak" or (
            kind == "html_inline" and re.fullmatch(r"<br\s*/?>", token.content, re.I)
        ):
            paragraph.add_run().add_break()
        else:
            raise ValueError(f"Unsupported inline Markdown: {kind} ({token.content!r})")


def configure_document(document, language, page_size):
    section = document.sections[0]
    if page_size == "a4":
        section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    else:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.7)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.04
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.widow_control = True
    normal.paragraph_format.keep_together = True
    lang = OxmlElement("w:lang")
    lang.set(qn("w:val"), "it-IT" if language == "it" else "en-GB")
    normal.element.get_or_add_rPr().append(lang)

    for name, size, before, after in [
        ("Title", 24, 0, 2), ("Subtitle", 12, 0, 4),
        ("Heading 1", 13, 12, 5), ("Heading 2", 11, 8, 3),
    ]:
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = name != "Subtitle"
        style.font.italic = False
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        # Clear decorative defaults inherited from python-docx's base template.
        for element in style.element.xpath("./w:pPr/w:pBdr | ./w:pPr/w:numPr | ./w:rPr/w:spacing"):
            element.getparent().remove(element)

    contact = document.styles.add_style("CV Contact", WD_STYLE_TYPE.PARAGRAPH)
    contact.base_style = normal
    contact.font.size = Pt(10)
    contact.paragraph_format.keep_with_next = True
    bullet = document.styles["List Bullet"]
    bullet.base_style = normal
    bullet.paragraph_format.left_indent = Inches(0.15)
    bullet.paragraph_format.first_line_indent = Inches(-0.15)
    bullet.paragraph_format.space_after = Pt(3)


def export_cv(source: Path, output: Path, *, photo=None, language="auto", page_size="letter"):
    text = source.read_text(encoding="utf-8-sig")
    if text.startswith("---\n"):
        parts = text.split("---\n", 2)
        if len(parts) != 3:
            raise ValueError("Unclosed YAML front matter")
        text = parts[2]
    if language == "auto":
        language = "it" if re.search(r"^## Profilo professionale", text, re.M | re.I) else "en"
    tokens = MARKDOWN.parse(text)
    document = Document()
    configure_document(document, language, page_size)
    document.core_properties.language = "it-IT" if language == "it" else "en-GB"
    document.core_properties.subject = "Curriculum Vitae"
    document.core_properties.comments = ""
    document.core_properties.last_modified_by = ""
    paragraph = None
    list_depth = 0
    subtitle_pending = False
    title_count = 0

    for index, token in enumerate(tokens):
        kind = token.type
        if kind == "heading_open":
            level = int(token.tag[1:])
            if level not in {1, 2, 3}:
                raise ValueError(f"Unsupported heading level: {level}")
            paragraph = document.add_paragraph(style="Title" if level == 1 else f"Heading {level - 1}")
            if level == 1:
                title_count += 1
                subtitle_pending = True
        elif kind == "paragraph_open":
            style = "List Bullet" if list_depth else "Subtitle" if subtitle_pending else "Normal"
            paragraph = document.add_paragraph(style=style)
            subtitle_pending = False
            children = tokens[index + 1].children or []
            if children and children[0].type == "code_inline":
                paragraph.paragraph_format.keep_with_next = True
                paragraph.paragraph_format.space_before = Pt(7)
        elif kind == "inline":
            if paragraph is None:
                raise ValueError("Inline text outside a paragraph")
            add_inline(paragraph, token.children)
        elif kind == "heading_close":
            if token.tag == "h1":
                name = paragraph.text
                document.core_properties.title = f"{name} - Curriculum Vitae"
                document.core_properties.author = name
                if photo:
                    section = document.sections[0]
                    width = section.page_width - section.left_margin - section.right_margin
                    paragraph.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
                    paragraph.add_run("\t").add_picture(str(photo), width=Inches(0.8))
                    picture = paragraph._p.xpath(".//wp:docPr")[-1]
                    picture.set("descr", f"Portrait of {name}" if language == "en" else f"Ritratto di {name}")
            paragraph = None
        elif kind == "paragraph_close":
            paragraph = None
        elif kind == "bullet_list_open":
            list_depth += 1
            if list_depth > 1:
                raise ValueError("Nested lists are not supported by the CV layout")
        elif kind == "bullet_list_close":
            list_depth -= 1
        elif kind in {"list_item_open", "list_item_close"}:
            continue
        elif kind == "html_block" and 'id="webaddress"' in token.content:
            contact = ContactParser()
            contact.feed(token.content)
            paragraph = document.add_paragraph(style="CV Contact")
            for content, href in contact.parts:
                add_text(paragraph, content, href=href)
            paragraph = None
        else:
            raise ValueError(f"Unsupported Markdown block: {kind}")
    if title_count != 1:
        raise ValueError("Expected exactly one level-one heading with the person's name")
    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    return output


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=ROOT / "index.md", help="Markdown CV (default: repository index.md)")
    parser.add_argument("--output", type=Path, help="DOCX output (default: <input directory>/build/cv-<language>.docx)")
    parser.add_argument("--language", choices=["auto", "en", "it"], default="auto")
    parser.add_argument("--page-size", choices=["letter", "a4"], default="letter")
    photos = parser.add_mutually_exclusive_group()
    photos.add_argument("--photo", type=Path, help="Custom local photo (default: profilo.jpeg beside input)")
    photos.add_argument("--no-photo", action="store_true", help="Omit the portrait")
    args = parser.parse_args(argv)
    source = args.input.resolve()
    try:
        text = source.read_text(encoding="utf-8-sig")
        language = args.language
        if language == "auto":
            language = "it" if re.search(r"^## Profilo professionale", text, re.M | re.I) else "en"
        photo = None if args.no_photo else args.photo or source.parent / "profilo.jpeg"
        if photo and not photo.is_file():
            raise ValueError(f"Photo not found: {photo}. Use --no-photo to omit it.")
        output = args.output.resolve() if args.output else source.parent / "build" / f"cv-{language}.docx"
        if output.suffix.lower() != ".docx":
            raise ValueError("Output must have the .docx extension")
        export_cv(source, output, photo=photo, language=language, page_size=args.page_size)
    except (OSError, ValueError) as error:
        parser.exit(1, f"Error: {error}\n")
    print(f"Created {output}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
