"""Content-preservation checks for the CV exporter (no office app required)."""

from html.parser import HTMLParser
from pathlib import Path
import re
import sys
import tempfile
import unittest

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from markdown_it import MarkdownIt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from export_docx import ROOT, export_cv, main


class VisibleText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data)


class ExportTests(unittest.TestCase):
    def test_project_content_links_and_photo_are_preserved(self):
        source = ROOT / "index.md"
        markdown = source.read_text(encoding="utf-8-sig").split("---\n", 2)[2]
        expected = VisibleText()
        expected.feed(MarkdownIt("commonmark", {"html": True}).render(markdown))
        with tempfile.TemporaryDirectory() as folder:
            output = Path(folder) / "cv.docx"
            export_cv(source, output, photo=ROOT / "profilo.jpeg")
            document = Document(output)
            actual = "".join(document.element.xpath(".//w:t/text()"))
            self.assertEqual(re.sub(r"\s", "", "".join(expected.parts)), re.sub(r"\s", "", actual))
            self.assertEqual(len(document.inline_shapes), 1)
            self.assertEqual(document.paragraphs[0].style.name, "Title")
            self.assertFalse(document.styles["Title"].element.xpath("./w:pPr/w:pBdr"))
            targets = {r.target_ref for r in document.part.rels.values() if r.reltype == RT.HYPERLINK}
            self.assertIn("mailto:p.alxzeta@gmail.com", targets)
            self.assertIn("tel:+393280169285", targets)
            self.assertIn("https://www.irion-edm.com/", targets)

    def test_italian_and_no_photo(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "index.md"
            source.write_text("# Nome Cognome\nRuolo\n\n## Profilo professionale\nCittà, qualità e C++.\n", encoding="utf-8")
            output = Path(folder) / "cv.docx"
            self.assertEqual(main(["--input", str(source), "--output", str(output), "--no-photo", "--page-size", "a4"]), 0)
            document = Document(output)
            self.assertEqual(document.core_properties.language, "it-IT")
            self.assertEqual(len(document.inline_shapes), 0)
            self.assertIn("Città, qualità e C++.", "".join(p.text for p in document.paragraphs))
            self.assertAlmostEqual(document.sections[0].page_width.inches, 8.2677, places=2)

    def test_unsupported_content_fails_without_creating_output(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "index.md"
            source.write_text("# Nome\n\n```python\nimportant_content()\n```\n", encoding="utf-8")
            output = Path(folder) / "cv.docx"
            with self.assertRaisesRegex(ValueError, "Unsupported Markdown"):
                export_cv(source, output)
            self.assertFalse(output.exists())

    def test_missing_photo_has_actionable_error(self):
        with tempfile.TemporaryDirectory() as folder:
            source = Path(folder) / "index.md"
            source.write_text("# Nome\n", encoding="utf-8")
            with self.assertRaises(SystemExit) as error:
                main(["--input", str(source)])
            self.assertEqual(error.exception.code, 1)
            self.assertFalse((Path(folder) / "build").exists())


if __name__ == "__main__":
    unittest.main()
